import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import copy

# Data for all 8 practices
PRACTICES = [
    {
        "nombre": "CONTROL DE DOS MOTORES DE CORRIENTE CONTINUA",
        "objetivo": "Controlar dos motores de corriente continua usando el L293D.",
        "material": [
            "1 Circuito integrado L293D (puente H).",
            "2 Motores de corriente continua (CD).",
            "1 Batería o fuente de alimentación de 5V.",
            "5 Estados lógicos (Logicstate en simulador para mandar 1 o 0).",
            "Terminales de alimentación (VCC/VSS) y tierra (GND)."
        ],
        "metodologia": (
            "Conectamos los 5 estados lógicos: los primeros 4 irán a cada una de las entradas del "
            "circuito integrado L293D (IN1, IN2, IN3 e IN4), y puenteamos los pines Enable 1 y Enable 2 "
            "para llevarlos a un quinto estado lógico y poder activar o desactivar ambos canales. "
            "Las salidas del L293D van conectadas en pares hacia los motores: las primeras dos (OUT1 y OUT2) "
            "para el primer motor y las otras dos (OUT3 y OUT4) para el segundo motor. "
            "Alimentamos las terminales VS y VSS a 5V y las tierras intermedias del integrado a GND. "
            "Finalmente, se simulan combinaciones lógicas en las entradas para verificar el giro a la derecha, "
            "a la izquierda y el frenado de cada motor."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Giro a la derecha del Motor 1 (sentido horario)",
                "desc": "Con el estado lógico de Enable en 1, colocamos IN1 en 1 e IN2 en 0. El motor 1 comienza a girar hacia la derecha (en el sentido de las manecillas del reloj).",
                "placeholder": "[Insertar captura de pantalla aquí: Motor 1 con giro a la derecha]"
            },
            {
                "titulo": "Movimiento 2: Giro a la izquierda del Motor 1 (sentido antihorario)",
                "desc": "Manteniendo Enable en 1, cambiamos la combinación a IN1 en 0 e IN2 en 1. La polaridad en las terminales se invierte y el motor 1 gira hacia la izquierda (en contra de las manecillas del reloj).",
                "placeholder": "[Insertar captura de pantalla aquí: Motor 1 con giro a la izquierda]"
            },
            {
                "titulo": "Movimiento 3: Control simultáneo e independiente del Motor 2",
                "desc": "Con Enable en 1, colocamos IN3 en 1 e IN4 en 0. El motor 2 gira hacia la derecha de forma completamente independiente al motor 1, verificando el funcionamiento del segundo puente H.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor 2 girando independientemente]"
            },
            {
                "titulo": "Movimiento 4: Paro o frenado del motor",
                "desc": "Al colocar ambas entradas en el mismo nivel lógico (IN1=0 e IN2=0, o ambas en 1), o al desactivar el Enable en 0, el motor frena y se detiene por completo.",
                "placeholder": "[Insertar captura de pantalla aquí: Motores detenidos en estado de paro]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí cómo funciona internamente el circuito integrado L293D como puente H y cómo mediante combinaciones lógicas de 1 y 0 podemos controlar el sentido de giro de dos motores de CD sin tener que invertir manualmente las conexiones.",
            "dificulto": "Al inicio se me dificultó identificar la distribución de pines del integrado porque tiene varios pines de tierra centrales y recordar qué combinación exacta de entradas hacía girar cada motor hacia la derecha o izquierda."
        }
    },
    {
        "nombre": "CONTROL MANUAL DE UN MOTOR UNIPOLAR SIN DRIVERS",
        "objetivo": "Controlar un motor unipolar por medio de un deep switch usando un ULN2803.",
        "material": [
            "1 Deep switch (dip switch de 4 posiciones).",
            "4 Resistencias de 10k ohms.",
            "1 Circuito integrado ULN2803 (arreglo Darlington).",
            "1 Motor unipolar de paso a paso.",
            "1 Fuente de poder de 5V y tierra (GND)."
        ],
        "metodologia": (
            "Conectamos el deep switch por la parte de ON con resistencias de 10k directamente hacia positivo "
            "y por la parte de OFF los puenteamos a tierra. Conectamos estas salidas del switch por la parte "
            "de ON hacia las entradas 1B, 2B, 3B y 4B del ULN2803 ordenadamente. "
            "La pata 10 (COM) del ULN2803 se conecta a la alimentación correspondiente. Posteriormente, "
            "las salidas 1C, 2C, 3C y 4C del ULN2803 se conectan directo a las entradas de las bobinas del "
            "motor unipolar. Las patas centrales de ambos lados (cables comunes del motor) las vamos a puentear "
            "y las conectamos a la línea positiva. De este modo, al cerrar cada interruptor del switch activamos "
            "manualmente una bobina a la vez."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Paso 1 (Activación de primera bobina a 0 grados)",
                "desc": "Se activa únicamente el interruptor 1 del switch (1B en alto y los demás en bajo). El rotor se alinea con la primera fase, quedando en la posición de referencia inicial (0°).",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar en paso 1 a 0 grados]"
            },
            {
                "titulo": "Movimiento 2: Paso 2 (Avance a 90 grados a la derecha)",
                "desc": "Se apaga el switch 1 y se enciende el switch 2. La corriente pasa a la segunda bobina y el eje del motor da un paso girando 90 grados hacia la derecha.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar con avance de 90 grados a la derecha]"
            },
            {
                "titulo": "Movimiento 3: Paso 3 y 4 (Avance a 180° y 270° a la derecha)",
                "desc": "Continuamos la secuencia manual activando el switch 3 y después el switch 4, completando los pasos hacia la derecha para realizar la rotación completa del eje.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar avanzando a 180 y 270 grados]"
            },
            {
                "titulo": "Movimiento 4: Giro en sentido antihorario (giro a la izquierda)",
                "desc": "Realizamos la secuencia de interruptores en orden inverso (switch 4, switch 3, switch 2 y switch 1), provocando que el eje del motor retroceda y gire hacia la izquierda.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar girando hacia la izquierda]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí cómo se energiza secuencialmente cada una de las fases de un motor unipolar para que avance paso a paso mecánicamente, y la función del ULN2803 para amplificar la corriente que demanda cada bobina.",
            "dificulto": "Lo que más se me dificultó fue coordinar el encendido y apagado de los interruptores del dip switch a mano en el orden exacto, ya que si me saltaba un paso o dejaba dos encendidos el motor se trababa o vibraba."
        }
    },
    {
        "nombre": "MOTOR UNIPOLAR CONTROL MANUAL OPCION B USANDO DRIVER L297",
        "objetivo": "Controlar un motor de forma manual usando la opción B por medio del driver L297.",
        "material": [
            "4 Estados lógicos (Logicstate para 0 y 1).",
            "1 Resistencia de 22k ohms.",
            "5 Resistencias de 1k ohms.",
            "1 Potenciómetro de 1k ohms.",
            "1 Capacitor de 3.3 uF.",
            "4 Transistores TIP122 (Darlington NPN).",
            "4 Diodos Zener.",
            "1 Motor unipolar paso a paso.",
            "1 Fuente de 5V y tierra.",
            "1 Controlador L297.",
            "1 Clock / pulsador manual de reloj."
        ],
        "metodologia": (
            "En esta práctica colocamos 4 estados lógicos conectados a las entradas 10 (ENABLE), 20 (RESET), "
            "19 (HALF/FULL) y 17 (CW/CCW) del L297: con el primero controlamos el enable, el pin 20 permite "
            "el reset y paso de datos, en la pata 19 controlamos el tamaño del paso (paso completo o medio paso) "
            "y en la 17 el sentido de giro. Conectamos alimentación y tierra; las patas HOME y CONTROL se puentean "
            "a tierra. VREF va al potenciómetro de 1k con su resistencia previa a positivo, y OSC se conecta a un "
            "capacitor de 3.3 uF a tierra y mediante una resistencia de 22k a positivo. Las terminales SENS1, SENS2 "
            "y SYNC van a tierra. Las salidas A, B, C y D del L297 pasan por resistencias de 1k cada una hacia la base "
            "de los 4 transistores TIP122. Los emisores van a tierra y los colectores se conectan a las bobinas del "
            "motor junto con los diodos Zener de protección. Los cables comunes del motor van a positivo. En la pata 18 "
            "(CLOCK) se aplica el pulso manual para hacer avanzar el motor paso a paso."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Giro a la derecha en paso completo (CW / 90 grados)",
                "desc": "Con el estado lógico CW/CCW en 1 y Half/Full en 0, cada pulso aplicado manualmente al pin de reloj hace avanzar el motor 90 grados hacia la derecha (paso completo).",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar girando a la derecha en paso completo (90°)]"
            },
            {
                "titulo": "Movimiento 2: Giro a la izquierda en paso completo (CCW / 90 grados)",
                "desc": "Cambiando el estado lógico de CW/CCW a 0, al aplicar los pulsos manuales el motor invierte la secuencia y gira hacia la izquierda a pasos de 90 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar girando a la izquierda en paso completo]"
            },
            {
                "titulo": "Movimiento 3: Giro en medio paso a 45 grados (Half Step)",
                "desc": "Cambiando el pin Half/Full a 1, cada pulso de reloj hace avanzar el motor en pasos más finos de 45 grados, permitiendo mayor resolución angular.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar girando en medio paso a 45 grados]"
            },
            {
                "titulo": "Movimiento 4: Estado de retención y paro",
                "desc": "Al dejar de enviar pulsos de reloj o mandar Enable a 0, el motor frena y se mantiene firmemente retenido en la última posición angular alcanzada.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor unipolar en estado de retención y paro]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí que el controlador L297 simplifica enormemente el control de motores paso a paso porque genera automáticamente la secuencia de fases correcta y nos permite conmutar entre paso completo y medio paso con solo cambiar un pin lógico.",
            "dificulto": "Se me dificultó comprender al inicio la red de oscilación (OSC) y el ajuste de referencia (VREF), así como cuidar que las conexiones de los 4 transistores TIP122 y los diodos Zener estuvieran exactamente orientados."
        }
    },
    {
        "nombre": "CONTROL AUTOMATICO DEL MOTOR UNIPOLAR",
        "objetivo": "Controlar de forma automática un motor usando un controlador L297 y una serie de compuertas AND.",
        "material": [
            "4 Estados lógicos para indicar 0 y 1.",
            "1 Resistencia de 4k ohms, 1 de 500k, 1 de 22k, 2 de 5k (o 0.5k) y 8 de 1k ohms.",
            "1 Potenciómetro de 1k ohms.",
            "Capacitores: 3.3 uF, tres de 1 uF, 10 uF (o 10 nF).",
            "4 Transistores TIP122.",
            "4 Diodos Zener.",
            "4 Compuertas lógicas AND.",
            "1 Circuito integrado temporizador NE555 (circuito generador de reloj).",
            "1 Controlador L297.",
            "1 Motor unipolar paso a paso.",
            "Fuente de 5V y tierra (GND)."
        ],
        "metodologia": (
            "Colocamos 4 estados lógicos conectados a las entradas 10, 20, 19 y 17 del L297 para controlar Enable, "
            "Reset, Half/Full y dirección CW/CCW. Las patas Home y Control se mandan a tierra. VREF se conecta al "
            "potenciómetro de 1k y OSC al capacitor de 3.3 uF y resistencia de 22k. Las salidas A, B, C y D del L297 "
            "pasan por resistencias de 1k y se conectan a una entrada de cada una de las 4 compuertas AND; las otras "
            "entradas de las compuertas se puentean en pares hacia las terminales INH1 e INH2 del L297. "
            "Las salidas de las compuertas AND van a la base de los 4 transistores TIP122 que controlan las bobinas del "
            "motor unipolar con sus diodos Zener. Para el funcionamiento automático, conectamos a la pata 18 (CLOCK) "
            "la salida (pin 3) de un circuito oscilador astable armado con un temporizador NE555, el cual incluye un "
            "potenciómetro de 500k que genera un tren continuo y regulable de pulsos."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Giro continuo automático a la derecha a 90 grados (4 pasos)",
                "desc": "Con CW/CCW en 1 y Half/Full en 0, el motor gira automáticamente de forma continua hacia la derecha dando 4 pasos de 90 grados por vuelta al ritmo del oscilador.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático a la derecha a 90 grados]"
            },
            {
                "titulo": "Movimiento 2: Giro continuo automático a la izquierda a 90 grados",
                "desc": "Cambiando CW/CCW a 0, el motor invierte automáticamente su sentido de rotación y continúa girando de manera continua hacia la izquierda a pasos de 90 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático a la izquierda a 90 grados]"
            },
            {
                "titulo": "Movimiento 3: Giro continuo automático en medio paso a 45 grados (8 pasos)",
                "desc": "Colocando el pin Half/Full en 1, el motor realiza automáticamente la secuencia de 8 medios pasos de 45 grados, observándose una rotación mucho más suave.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático en medio paso a 45 grados (8 pasos)]"
            },
            {
                "titulo": "Movimiento 4: Regulación de velocidad mediante el potenciómetro del NE555",
                "desc": "Al ajustar el potenciómetro de 500k en el circuito del 555, cambia la frecuencia de los pulsos del reloj, haciendo que el motor aumente o disminuya su velocidad de giro continuo en tiempo real.",
                "placeholder": "[Insertar captura de pantalla aquí: Variación de velocidad de giro con potenciómetro]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí a integrar un temporizador NE555 en modo astable junto con el L297 y compuertas AND para lograr un funcionamiento totalmente automático y continuo del motor, entendiendo cómo la frecuencia de reloj determina la velocidad angular.",
            "dificulto": "Lo más complejo fue realizar todo el cableado en el simulador sin equivocarse, ya que conectar el 555, las compuertas AND, el L297 y los transistores implicaba muchas líneas cruzadas."
        }
    },
    {
        "nombre": "BIPOLAR TOTALMENTE MANUAL MOTOR PASO A PASO",
        "objetivo": "Controlar un motor paso a paso con una serie de combinaciones con botones otorgando valores de 0s o 1s.",
        "material": [
            "4 Estados lógicos (Logicstate) para las entradas y 1 para Enable.",
            "1 Circuito integrado driver L293D.",
            "1 Motor bipolar de paso a paso (4 hilos / 2 bobinas).",
            "4 Diodos 1N4148 de protección.",
            "Batería o fuente de poder de 5V.",
            "Conexión a tierra (GND)."
        ],
        "metodologia": (
            "Conectamos los estados lógicos a las 4 entradas IN1, IN2, IN3 e IN4 del L293D, y otro estado lógico "
            "conectado a los dos puertos Enable puenteados entre sí. Conectamos VSS y VS a positivo (+5V) y los "
            "pines GND centrales a tierra. Las 4 salidas del L293D se conectan por pares a cada bobina del motor "
            "bipolar: la bobina izquierda recibe el par de salidas 1 y 2, y la bobina derecha recibe el par de salidas "
            "3 y 4. Por cada salida colocamos un diodo 1N4148 con el cátodo hacia la alimentación positiva como "
            "protección contra sobretensiones. Para hacer girar el motor, se introducen manualmente las combinaciones "
            "lógicas que van alternando la polaridad de ambas bobinas paso a paso."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Paso 1 (Posición inicial a 0 grados)",
                "desc": "Se coloca la primera combinación lógica de la secuencia bipolar (ej. IN1=1, IN2=0, IN3=1, IN4=0 con Enable=1). El rotor se alinea con la primera fase en la posición de referencia a 0 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar en posición inicial de 0 grados]"
            },
            {
                "titulo": "Movimiento 2: Paso 2 (Giro a la derecha a 90 grados)",
                "desc": "Se cambia a la siguiente combinación (ej. IN1=0, IN2=1, IN3=1, IN4=0), invirtiendo la polaridad de la primera bobina. El eje da un paso girando 90 grados en sentido horario.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar girando a la derecha a 90 grados]"
            },
            {
                "titulo": "Movimiento 3: Paso 3 y 4 (Avance a 180° y 270° en sentido horario)",
                "desc": "Se aplican las combinaciones 3 y 4 de la secuencia bipolar, completando la vuelta completa hacia la derecha mediante la conmutación manual de las fases.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar avanzando a 180 y 270 grados]"
            },
            {
                "titulo": "Movimiento 4: Giro en sentido antihorario (giro a la izquierda)",
                "desc": "Se introducen los estados lógicos en orden inverso al anterior, haciendo que el flujo de corriente circule en sentido contrario y el motor gire hacia la izquierda.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar girando a la izquierda en sentido antihorario]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí que en un motor bipolar es indispensable invertir la polaridad de la corriente en las bobinas hacia adelante y hacia atrás para generar los pasos, y por eso el L293D es ideal al funcionar como puente H doble.",
            "dificulto": "Se me dificultó aprender de memoria y conmutar a mano la tabla de combinaciones lógicas, ya que equivocarse en un solo estado provoca que el motor no avance o regrese al paso anterior."
        }
    },
    {
        "nombre": "MOTOR BIPOLAR CONTROL MANUAL",
        "objetivo": "Controlar un motor bipolar de forma manual.",
        "material": [
            "4 Estados lógicos para indicar 0 y 1.",
            "Resistencias: 4k, 500k, 22k, dos de 5k y ocho de 1k ohms.",
            "1 Potenciómetro de 1k ohms.",
            "Capacitores: 3.3 uF, tres de 1 uF, 10 uF, 7 uF y dos de 100 uF.",
            "1 Controlador L297.",
            "1 Driver de potencia puente H L298.",
            "Puente de 8 diodos de protección rápida.",
            "1 Motor bipolar paso a paso.",
            "Fuente de 5V y tierra (GND).",
            "Circuito de pulso de reloj manual."
        ],
        "metodologia": (
            "Conectamos 4 estados lógicos a las entradas 10 (ENABLE), 20 (RESET), 19 (HALF/FULL) y 17 (CW/CCW) "
            "del L297 ordenadamente para controlar el estado, tamaño de paso y sentido de giro. Conectamos "
            "alimentación y tierra; Home y Control se puentean a tierra. VREF se conecta al potenciómetro de 1k con su "
            "resistencia previa, y OSC al capacitor de 3.3 uF a tierra y resistencia de 22k a positivo. "
            "Las salidas A, B, C y D del L297 se conectan directamente a las 4 entradas del driver L298, y las señales "
            "INH1 e INH2 se llevan a los enables del L298. Las salidas del L298 van conectadas a las dos bobinas del "
            "motor bipolar protegidas mediante un puente de 8 diodos conectados a positivo y tierra. "
            "Al aplicar pulsos manuales en la pata 18 (CLOCK), el motor avanza un paso por cada pulso ingresado."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Giro a la derecha a 90 grados (Paso completo - CW)",
                "desc": "Con CW/CCW en 1 y Half/Full en 0, cada pulso manual enviado al pin de reloj hace girar el motor 90 grados hacia la derecha (paso completo).",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar en giro a la derecha a 90 grados]"
            },
            {
                "titulo": "Movimiento 2: Giro a la izquierda a 90 grados (Paso completo - CCW)",
                "desc": "Cambiando el estado lógico de CW/CCW a 0, cada pulso manual hace avanzar el motor 90 grados hacia la izquierda en sentido antihorario.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar en giro a la izquierda a 90 grados]"
            },
            {
                "titulo": "Movimiento 3: Giro en medio paso a 45 grados (Half Step)",
                "desc": "Cambiando Half/Full a 1, cada pulso manual produce un desplazamiento más fino de 45 grados por paso, permitiendo el doble de paradas por vuelta.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar en medio paso a 45 grados]"
            },
            {
                "titulo": "Movimiento 4: Retención y parada del motor",
                "desc": "Al dejar de enviar pulsos o deshabilitar Enable, el motor frena y se mantiene firmemente retenido en la posición actual.",
                "placeholder": "[Insertar captura de pantalla aquí: Motor bipolar en estado de parada y retención]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí a utilizar la combinación del integrado L297 como generador de secuencia lógica junto con el L298 como puente de potencia para controlar motores bipolares sin necesidad de microcontrolador.",
            "dificulto": "Se me dificultó el cableado del puente de 8 diodos de protección externa y las conexiones entre el L297 y el L298, debido a la gran cantidad de pines que deben coincidir exactamente."
        }
    },
    {
        "nombre": "MOTOR BIPOLAR CONTROL AUTOMATICO",
        "objetivo": "Controlar un motor bipolar de forma automatica.",
        "material": [
            "4 Estados lógicos para indicar 0 y 1.",
            "Resistencias: 4k, 500k, 22k, dos de 5k y ocho de 1k ohms.",
            "1 Potenciómetro de 1k y 1 potenciómetro de 500k ohms.",
            "Capacitores: 3.3 uF, tres de 1 uF, 10 uF, 7 uF, dos de 100 uF y 10 nF.",
            "1 Controlador L297.",
            "1 Driver de potencia L298.",
            "1 Circuito temporizador NE555 (circuito oscilador de reloj).",
            "Puente de 8 diodos de protección.",
            "1 Motor bipolar paso a paso.",
            "Fuente de poder de 5V y tierra (GND)."
        ],
        "metodologia": (
            "Se implementa el circuito base del controlador L297 conectado al driver L298 y las salidas hacia "
            "las bobinas del motor bipolar con su puente de 8 diodos. Los estados lógicos en Half/Full y CW/CCW "
            "definen el modo de paso y el sentido de giro. Para lograr el control automático, en la pata 18 (CLOCK) "
            "del L297 se conecta la señal continua de reloj generada por un circuito integrado NE555 en configuración "
            "astable. El circuito del 555 cuenta con un potenciómetro de 500k que permite variar la frecuencia de los "
            "pulsos, logrando que el motor gire de forma continua y automatizada a la velocidad deseada."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Giro continuo automático a la derecha a 90 grados (4 pasos de 90°)",
                "desc": "Con CW/CCW en 1 y Half/Full en 0, el motor gira automáticamente de forma continua hacia la derecha dando 4 pasos de 90 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático a la derecha a 90 grados (4 pasos)]"
            },
            {
                "titulo": "Movimiento 2: Giro continuo automático a la izquierda a 90 grados (4 pasos de 90°)",
                "desc": "Cambiando el pin CW/CCW a 0, el motor invierte automáticamente su dirección y gira de forma continua hacia la izquierda a pasos de 90 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático a la izquierda a 90 grados (4 pasos)]"
            },
            {
                "titulo": "Movimiento 3: Giro continuo automático en medio paso a 45 grados (8 pasos de 45°)",
                "desc": "Colocando Half/Full en 1, el motor conmuta automáticamente en 8 medios pasos de 45 grados por ciclo, produciendo una rotación más uniforme y suave.",
                "placeholder": "[Insertar captura de pantalla aquí: Giro continuo automático en medio paso a 45 grados (8 pasos)]"
            },
            {
                "titulo": "Movimiento 4: Ajuste de velocidad de giro mediante el potenciómetro",
                "desc": "Al girar el potenciómetro de 500k del circuito NE555, cambia la frecuencia de los pulsos generados, aumentando o disminuyendo la velocidad de rotación continua del motor en tiempo real.",
                "placeholder": "[Insertar captura de pantalla aquí: Variación de velocidad de giro con el potenciómetro]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí a automatizar completamente la rotación de un motor bipolar usando un oscilador 555 para suministrar el tren de pulsos de reloj al L297, comprendiendo el efecto de la frecuencia en la velocidad del motor.",
            "dificulto": "Se me dificultó calibrar el potenciómetro y los capacitores del NE555 en la simulación, ya que si la frecuencia era demasiado alta el motor perdía pasos y se frenaba."
        }
    },
    {
        "nombre": "SERVO MOTOR",
        "objetivo": "Controlar o verificar el funcionamiento de un servo motor.",
        "material": [
            "1 Resistencia de 220k ohms.",
            "1 Potenciómetro de 10k ohms.",
            "1 Resistencia de 15k ohms.",
            "2 Resistencias de 10k ohms.",
            "2 Capacitores de 100 uF (o cerámicos correspondientes).",
            "1 Transistor 2N3904 (o 2N1711).",
            "1 Fuente de alimentación de 6 voltios.",
            "1 Servomotor (con control por ancho de pulso PWM).",
            "1 Circuito integrado temporizador NE555."
        ],
        "metodologia": (
            "Comenzamos por la fuente en la parte positiva, esta baja a una resistencia de 220k, después va a un "
            "potenciómetro de 10k y por el otro extremo baja a una resistencia de 15k para después tener un capacitor "
            "conectado a tierra. La terminal central del potenciómetro irá al pin 7 (DISCHARGE) del NE555. "
            "Las patas 6 y 2 (THRESHOLD y TRIGGER) estarán puenteadas entre la resistencia de 15k y el capacitor. "
            "GND va a negativo y VCC a positivo (+6V). La pata 5 se conecta a un capacitor que va a negativo. "
            "En la pata 3 (OUTPUT) conectamos una resistencia de 10k que va a la base del transistor; el colector "
            "del transistor tiene una resistencia de 10k a positivo y se conecta al cable de señal de control del "
            "servomotor. El emisor del transistor va a negativo. Finalizamos conectando el servomotor a positivo (+6V) "
            "y negativo (GND). Al ajustar el potenciómetro varía la duración del pulso (de 1 ms a 2 ms), posicionando "
            "el eje del servomotor en el ángulo deseado."
        ),
        "resultados": [
            {
                "titulo": "Movimiento 1: Posición a 0 grados / Extremo izquierdo (Pulso de ~1 ms)",
                "desc": "Con el potenciómetro en su valor mínimo, el circuito entrega un pulso de aproximadamente 1 ms de duración, haciendo que el servomotor se posicione en el extremo de 0 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Servomotor en posición de 0 grados]"
            },
            {
                "titulo": "Movimiento 2: Posición a 90 grados / Posición central (Pulso de ~1.5 ms)",
                "desc": "Ajustando el potenciómetro al 50% de su recorrido, el ancho de pulso se sitúa en aproximadamente 1.5 ms, colocando el servomotor exactamente en su posición central de 90 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Servomotor en posición central a 90 grados (al 50%)]"
            },
            {
                "titulo": "Movimiento 3: Posición a 180 grados / Extremo derecho (Pulso de ~2 ms)",
                "desc": "Al colocar el potenciómetro al 100% de su recorrido, el ancho de pulso alcanza aproximadamente 2 ms, desplazando el servomotor a su posición límite de 180 grados.",
                "placeholder": "[Insertar captura de pantalla aquí: Servomotor en posición de 180 grados (al 100%)]"
            },
            {
                "titulo": "Movimiento 4: Movimiento angular continuo y progresivo",
                "desc": "Al girar lentamente el potenciómetro entre sus extremos, el servomotor realiza un barrido angular suave y continuo respondiendo en tiempo real al cambio de ancho de pulso.",
                "placeholder": "[Insertar captura de pantalla aquí: Barrido angular continuo del servomotor de 0 a 180 grados]"
            }
        ],
        "observaciones": {
            "aprendio": "Aprendí el funcionamiento de los servomotores mediante modulación por ancho de pulso (PWM) y cómo un integrado 555 con una etapa de transistor puede modular pulsos de 1 ms a 2 ms cada 20 ms para controlar con precisión el ángulo del eje.",
            "dificulto": "Se me dificultó comprender la diferencia entre la frecuencia total del pulso y el tiempo en alto (duty cycle), además de calibrar el potenciómetro para que el recorrido cubriera de 0° a 180° completos."
        }
    }
]

print(f"Loaded {len(PRACTICES)} practices configurations successfully.")

def format_run(run, text, bold=False, italic=False, size_pt=10.5, font_name="Comfortaa"):
    run.text = text
    run.bold = bold
    run.italic = italic
    if size_pt:
        run.font.size = Pt(size_pt)
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def add_para(cell, text="", bold=False, italic=False, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.15):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        format_run(r, text, bold=bold, italic=italic, size_pt=size_pt)
    return p

def populate_cell(cell, practice):
    # Clear existing paragraphs in cell
    p_elements = cell._tc.xpath('./w:p')
    for p_elem in p_elements:
        cell._tc.remove(p_elem)
        
    # P0: Nombre del alumno
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Nombre del alumno: ")
    format_run(r_label, "Nombre del alumno: ", bold=False, size_pt=10.5)
    r_val = p.add_run("Altamirano Plantillas Eddie David")
    format_run(r_val, "Altamirano Plantillas Eddie David", bold=True, size_pt=10.5)

    # P1: Nombre de la Práctica
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Nombre de la Práctica : ")
    format_run(r_label, "Nombre de la Práctica : ", bold=True, size_pt=11)
    r_val = p.add_run(practice["nombre"])
    format_run(r_val, practice["nombre"], bold=True, size_pt=11)

    # P2: Objetivo
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Objetivo:")
    format_run(r_label, "Objetivo:", bold=True, size_pt=10.5)
    
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_val = p.add_run(practice["objetivo"])
    format_run(r_val, practice["objetivo"], bold=False, size_pt=10.5)

    # P3: Material y Equipo
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Material y Equipo:")
    format_run(r_label, "Material y Equipo:", bold=True, size_pt=10.5)

    for item in practice["material"]:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r_bullet = p.add_run("•  ")
        format_run(r_bullet, "•  ", bold=True, size_pt=10)
        r_item = p.add_run(item)
        format_run(r_item, item, bold=False, size_pt=10.5)

    # P4: Metodología
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Metodología: (INDICA LOS PASOS PARA REALIZAR LA PRÁCTICA)")
    format_run(r_label, "Metodología: (INDICA LOS PASOS PARA REALIZAR LA PRÁCTICA)", bold=True, size_pt=10.5)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r_val = p.add_run(practice["metodologia"])
    format_run(r_val, practice["metodologia"], bold=False, size_pt=10.5)

    # P5: Resultados
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Resultados (IMÁGENES, RESPUESTA DE LAS PREGUNTAS)")
    format_run(r_label, "Resultados (IMÁGENES, RESPUESTA DE LAS PREGUNTAS)", bold=True, size_pt=10.5)

    for res in practice["resultados"]:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r_title = p.add_run(res["titulo"])
        format_run(r_title, res["titulo"], bold=True, size_pt=10.5)

        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_desc = p.add_run(res["desc"])
        format_run(r_desc, res["desc"], bold=False, size_pt=10)

        # Placeholder for image
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_box = p.add_run(f"   {res['placeholder']}")
        format_run(r_box, f"   {res['placeholder']}", bold=False, italic=True, size_pt=9.5)

    # P6: Observaciones y comentarios
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Observaciones y comentarios: (QUE APRENDIÓ, QUE SE LE DIFICULTÓ)")
    format_run(r_label, "Observaciones y comentarios: (QUE APRENDIÓ, QUE SE LE DIFICULTÓ)", bold=True, size_pt=10.5)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r_ap_label = p.add_run("Qué aprendió: ")
    format_run(r_ap_label, "Qué aprendió: ", bold=True, size_pt=10.5)
    r_ap_val = p.add_run(practice["observaciones"]["aprendio"])
    format_run(r_ap_val, practice["observaciones"]["aprendio"], bold=False, size_pt=10.5)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_dif_label = p.add_run("Qué se le dificultó: ")
    format_run(r_dif_label, "Qué se le dificultó: ", bold=True, size_pt=10.5)
    r_dif_val = p.add_run(practice["observaciones"]["dificulto"])
    format_run(r_dif_val, practice["observaciones"]["dificulto"], bold=False, size_pt=10.5)

    # P7: Fuentes de Información
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_label = p.add_run("Fuentes de Información")
    format_run(r_label, "Fuentes de Información", bold=True, size_pt=10.5)


def build_full_document():
    # Load backup template
    doc = docx.Document('FORMATO PRACTICAS SP_backup.docx')
    base_tbl = doc.tables[0]._tbl
    
    # Store header paragraphs XML
    header_p_xmls = [p._p for p in doc.paragraphs[:5]]
    
    # Clear the initial document body except style definitions
    # Actually, we can populate the first practice in the existing table!
    populate_cell(doc.tables[0].rows[2].cells[0], PRACTICES[0])
    
    # Now for practices 1 to 7 (practices 2 to 8):
    for i in range(1, len(PRACTICES)):
        # Add page break
        doc.add_page_break()
        
        # Add header paragraphs
        for hp in [
            ("INSTITUTO TECNOLÓGICO DE TEPIC", True),
            ("INGENIERÍA EN SISTEMAS COMPUTACIONALES", False),
            ("Sistemas Programables", False),
            ("Reporte de Práctica", False),
            ("", False)
        ]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if hp[0] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            if hp[0]:
                r = p.add_run(hp[0])
                format_run(r, hp[0], bold=hp[1], size_pt=11)
        
        # Deepcopy table XML and append
        new_tbl_xml = copy.deepcopy(base_tbl)
        doc._body._body.append(new_tbl_xml)
        
        # Get the new table from doc.tables
        current_tbl = doc.tables[-1]
        
        # Populate Row 2 of current table
        populate_cell(current_tbl.rows[2].cells[0], PRACTICES[i])
        
    doc.save('FORMATO PRACTICAS SP.docx')
    print("Successfully generated FORMATO PRACTICAS SP.docx with all 8 practices!")

if __name__ == '__main__':
    build_full_document()

